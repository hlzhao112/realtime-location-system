import unittest
from io import BytesIO

from openpyxl import Workbook

from app.services.excel_io import (
    TZ_ALIASES,
    TZ_HEADERS,
    _map_headers,
    build_lines_from_pos,
    parse_position_file,
    parse_trolley_file,
    read_first_sheet,
)
from app.services.normalize import norm_reader_id


def _xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class MappingImportTest(unittest.TestCase):
    def test_reader_keeps_mixed_codes(self):
        self.assertEqual(norm_reader_id("LYG01"), "LYG01")
        self.assertEqual(norm_reader_id("lyg16"), "lyg16")
        self.assertEqual(norm_reader_id("RFID1"), "RFID1")
        self.assertEqual(norm_reader_id("OF1"), "OF1")
        self.assertEqual(norm_reader_id("  RFID2  "), "RFID2")

    def test_tz_header_accepts_seat_install(self):
        idx = _map_headers(["台座名称", "编号", "台座安装"], TZ_ALIASES, TZ_HEADERS)
        self.assertEqual(idx["台座安装LYG"], 2)

    def test_parse_trolley_keeps_reader_codes(self):
        data = _xlsx(
            [
                ["台座名称", "编号", "台座安装"],
                ["台座编号1-1", "TZ_2796", "RFID1"],
                ["台座编号1-2", "TZ_2797", "OF1"],
                ["台座编号1-3", "TZ_2798", "LYG01"],
            ]
        )
        rows, errs = parse_trolley_file(data)
        self.assertFalse(errs)
        self.assertEqual([r["reader"] for r in rows], ["RFID1", "OF1", "LYG01"])

    def test_last_row_not_dropped(self):
        data = _xlsx(
            [
                ["台座名称", "编号", "台座安装"],
                ["A", "TZ_1", "LYG1"],
                ["B", "TZ_2", "LYG2"],
                ["最后一台", "TZ_LAST", "LYG99"],
            ]
        )
        sheet = read_first_sheet(data)
        self.assertEqual(len(sheet), 4)
        rows, errs = parse_trolley_file(data)
        self.assertFalse(errs)
        self.assertEqual(rows[-1]["tz"], "TZ_LAST")
        self.assertEqual(rows[-1]["name"], "最后一台")

    def test_same_proc_code_keeps_both_names(self):
        pos = [
            {
                "生产线": "1#产线",
                "生产线ID": "ZCX_1",
                "工序": "蒸养",
                "工序ID": "GX_0010",
                "区域": "蒸养区",
                "区域ID": "QY_1",
                "蓝牙标签编号": "B1",
            },
            {
                "生产线": "1#产线",
                "生产线ID": "ZCX_1",
                "工序": "自然养护",
                "工序ID": "GX_0010",
                "区域": "养护区",
                "区域ID": "QY_2",
                "蓝牙标签编号": "B2",
            },
        ]
        lines, used, undefined = build_lines_from_pos(pos, {"B1": ["E1"], "B2": ["E2"]})
        names = [p["name"] for p in lines[0]["procs"]]
        self.assertEqual(names, ["蒸养", "自然养护"])
        self.assertEqual(lines[0]["procs"][0]["code"], "GX_0010")
        self.assertEqual(lines[0]["procs"][1]["code"], "GX_0010")
        self.assertEqual(undefined, [])

    def test_position_carry_keeps_last_area_tags(self):
        data = _xlsx(
            [
                ["生产线", "生产线ID", "工序", "工序ID", "区域", "区域ID", "蓝牙标签编号"],
                ["1#产线", "ZCX_1", "蒸养", "GX_1", "蒸养区", "QY_1", "B1"],
                ["", "", "", "", "", "", "B2"],
                ["", "", "", "", "养护区", "QY_2", "B3"],
                ["", "", "", "", "", "", "B4"],
            ]
        )
        rows, errs = parse_position_file(data)
        self.assertFalse(any(e.get("level") == "err" for e in errs))
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[-1]["区域ID"], "QY_2")
        self.assertEqual(rows[-1]["蓝牙标签编号"], "B4")
        lines, used, undefined = build_lines_from_pos(rows, {"B1": ["E1"], "B2": ["E2"], "B3": ["E3"], "B4": ["E4"]})
        areas = lines[0]["procs"][0]["areas"]
        self.assertEqual([a["id"] for a in areas], ["QY_1", "QY_2"])
        self.assertEqual(areas[-1]["nos"], ["B3", "B4"])
        self.assertEqual(undefined, [])

    def test_position_merged_cells_keep_last_area(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["生产线", "生产线ID", "工序", "工序ID", "区域", "区域ID", "蓝牙标签编号"])
        ws.append(["1#产线", "ZCX_1", "蒸养", "GX_1", "蒸养区", "QY_1", "B1"])
        ws.append([None, None, None, None, "养护区", "QY_2", "B2"])
        ws.merge_cells("A2:A3")
        ws.merge_cells("B2:B3")
        ws.merge_cells("C2:C3")
        ws.merge_cells("D2:D3")
        buf = BytesIO()
        wb.save(buf)
        rows, errs = parse_position_file(buf.getvalue())
        self.assertFalse(any(e.get("level") == "err" for e in errs))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[-1]["区域"], "养护区")
        self.assertEqual(rows[-1]["蓝牙标签编号"], "B2")

    def test_replace_mapping_persists_last_tag_epcs(self):
        from sqlalchemy import create_engine, func, select
        from sqlalchemy.orm import sessionmaker

        from app.db import Base
        from app.models import Epc, Project, TagNo
        from app.services.project_svc import replace_mapping

        eng = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(eng)
        Session = sessionmaker(bind=eng, autoflush=False)
        db = Session()
        p = Project(pid="PYT", name="t", ingest_token="tok-last-epc")
        db.add(p)
        db.flush()
        lines = [
            {
                "id": "ZCX_1",
                "name": "1#",
                "procs": [
                    {
                        "code": "GX_1",
                        "name": "压浆",
                        "areas": [{"id": "QY_1", "name": "区", "nos": ["B142", "B143", "B144"], "epcs": []}],
                    }
                ],
            }
        ]
        tag_nos = {
            "B142": ["E2 00 01"],
            "B143": ["E2 00 02"],
            "B144": ["E2 00 03", "E2 00 04", "E2 00 05"],
        }
        replace_mapping(db, p, lines, tag_nos)
        db.commit()
        db2 = Session()
        epc_n = dict(
            db2.execute(select(TagNo.no, func.count(Epc.id)).outerjoin(Epc).group_by(TagNo.no)).all()
        )
        self.assertEqual(epc_n["B142"], 1)
        self.assertEqual(epc_n["B143"], 1)
        self.assertEqual(epc_n["B144"], 3)


if __name__ == "__main__":
    unittest.main()
